from pathlib import Path
import math, re, textwrap
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).parent
A=ROOT/'assets'; A.mkdir(exist_ok=True)
INK='#182d43'; BLUE='#eaf0f6'; RED='#8e1b25'; GRAY='#536171'
def font(size,bold=False): return ImageFont.truetype('C:/Windows/Fonts/'+('arialbd.ttf' if bold else 'arial.ttf'),size)
def canvas(w,h):
    im=Image.new('RGB',(w,h),'white'); return im,ImageDraw.Draw(im)
def label(d,xy,text,size=30,width=35,fill=INK):
    lines=[]
    for p in text.split('\n'): lines.extend(textwrap.wrap(p,width=width) or [''])
    text='\n'.join(lines)
    d.multiline_text(xy,text,font=font(size),fill=fill,anchor='mm',align='center',spacing=8)
def box(d,x,y,w,h,text,kind='box',size=30):
    bounds=(x-w/2,y-h/2,x+w/2,y+h/2)
    if kind=='diamond': d.polygon([(x,y-h/2),(x+w/2,y),(x,y+h/2),(x-w/2,y)],fill=BLUE,outline=INK,width=3)
    elif kind=='oval': d.ellipse(bounds,fill=BLUE,outline=INK,width=3)
    else: d.rounded_rectangle(bounds,radius=16,fill=BLUE,outline=INK,width=3)
    label(d,(x,y),text,size,max(12,int(w/(size*.57))-2))
def arrow(d,points,text=None,at=None,triangle=False):
    d.line(points,fill=GRAY,width=3)
    x,y=points[-1]; px,py=points[-2]; a=math.atan2(y-py,x-px)
    tri=[(x,y),(x-17*math.cos(a-.45),y-17*math.sin(a-.45)),(x-17*math.cos(a+.45),y-17*math.sin(a+.45))]
    d.polygon(tri,fill='white' if triangle else GRAY,outline=GRAY,width=3)
    if text: label(d,at,text,25,30)
def actor(d,x,y,name):
    d.ellipse((x-21,y-65,x+21,y-23),outline=INK,width=4)
    d.line([(x,y-23),(x,y+45)],fill=INK,width=4)
    d.line([(x-42,y),(x+42,y)],fill=INK,width=4)
    d.line([(x-35,y+90),(x,y+45),(x+35,y+90)],fill=INK,width=4)
    label(d,(x,y+130),name,27,20)

im,d=canvas(1450,1720)
label(d,(725,50),'PROVISIONAL EXISTING PROCESS',36,60)
label(d,(725,100),'Validate each step with library personnel.',27,70)
nodes=[(530,210,560,90,'Borrower requests or searches for a book','oval'),(530,365,560,100,'Staff check records and shelf','box'),(530,560,470,200,'Copy available?','diamond'),(530,795,560,110,'Verify borrower and applicable policy','box'),(530,995,560,110,'Record loan and due date; release book','box'),(530,1195,560,110,'Receive return; inspect and update records','box'),(530,1395,560,110,'Review outstanding loans and prepare reports','box'),(530,1580,360,80,'End','oval')]
for n in nodes: box(d,*n)
for top,bot in [(255,315),(415,460),(660,740),(850,940),(1050,1140),(1250,1340),(1450,1540)]: arrow(d,[(530,top),(530,bot)])
label(d,(585,700),'Yes',25)
box(d,1135,790,420,150,'Inform borrower; record request only if existing policy permits',size=28)
arrow(d,[(765,560),(1135,560),(1135,715)],'No',(965,525))
box(d,1135,1030,360,80,'End inquiry','oval')
arrow(d,[(1135,865),(1135,990)])
im.save(A/'existing_flow.png')

exec((ROOT/'diagrams_v2.py').read_text(encoding='utf-8'))

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Inches(.85); sec.bottom_margin=Inches(.85)
sec.left_margin=Inches(1.15); sec.right_margin=Inches(.85)
sec.header_distance=Inches(.35); sec.footer_distance=Inches(.35)
normal=doc.styles['Normal']; normal.font.name='Times New Roman'; normal.font.size=Pt(12)
normal.paragraph_format.line_spacing=1.5; normal.paragraph_format.space_after=Pt(6)
for key,size in [('Title',20),('Heading 1',16),('Heading 2',13),('Heading 3',12)]:
    s=doc.styles[key]; s.font.name='Times New Roman'; s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string('182D43')
    s.paragraph_format.keep_with_next=True
    s.paragraph_format.space_before=Pt(14); s.paragraph_format.space_after=Pt(8)
doc.styles['Caption'].font.name='Times New Roman'; doc.styles['Caption'].font.size=Pt(10)
doc.styles['Caption'].paragraph_format.line_spacing=1.1
sec.different_first_page_header_footer=True
h=sec.header.paragraphs[0]; h.text='IBA COLLEGE OF MINDANAO, INC.  |  LIBRARY SYSTEM'; h.style='Caption'
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
f.add_run('Page ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)

lines=(ROOT/'IBA_Library_System_REVISED.md').read_text(encoding='utf-8').splitlines()
i=0; cover=True
while i<len(lines):
    line=lines[i].strip(); i+=1
    if not line: continue
    if line.startswith('# CHAPTER'):
        cover=False
        p=doc.add_heading(line[2:],1); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.page_break_before=True
    elif line.startswith('# '):
        p=doc.add_paragraph(line[2:],'Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after=Pt(38)
    elif line.startswith('### '): doc.add_heading(line[4:],3)
    elif line.startswith('## '): doc.add_heading(line[3:],2)
    elif line.startswith('!['):
        m=re.match(r'!\[(.*?)\]\((.*?)\)',line); path=ROOT/m[2]
        if not path.exists(): raise RuntimeError('Missing figure '+str(path))
        w,h=Image.open(path).size; width=min(6.3,7.1*w/h)
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(path),width=Inches(width))
        p=doc.add_paragraph(m[1],'Caption'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    elif line.startswith('|'):
        rows=[line]
        while i<len(lines) and lines[i].strip().startswith('|'):
            rows.append(lines[i].strip()); i+=1
        rows=[r for r in rows if not re.match(r'^\|[\s:|-]+\|$',r)]
        data=[[c.strip() for c in r.strip('|').split('|')] for r in rows]
        t=doc.add_table(rows=0,cols=len(data[0])); t.style='Table Grid'
        for ridx,row in enumerate(data):
            cells=t.add_row().cells
            for c,text in zip(cells,row):
                c.text=text
                for p in c.paragraphs:
                    p.paragraph_format.line_spacing=1.1; p.paragraph_format.space_after=Pt(5)
                    if ridx==0: p.paragraph_format.keep_with_next=True
                    for r in p.runs: r.font.size=Pt(10); r.bold=(ridx==0)
                if ridx==0:
                    shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'EAF0F6'); c._tc.get_or_add_tcPr().append(shade)
            trPr=t.rows[-1]._tr.get_or_add_trPr()
            no_split=OxmlElement('w:cantSplit'); trPr.append(no_split)
            if ridx==0: trPr.append(OxmlElement('w:tblHeader'))
        doc.add_paragraph().paragraph_format.space_after=Pt(0)
    else:
        p=doc.add_paragraph(line)
        if cover:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after=Pt(14)
        elif not re.match(r'^\d+\.',line): p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY

doc.core_properties.title='Library Reservation and Circulation Management System — Revised Paper'
doc.core_properties.author='Ortiz; Tavita; Dalayon; Garcia'
doc.core_properties.subject='Six-chapter paper aligned to the local system implementation'
dest=ROOT/'IBA_Library_System_REVISED.docx'; doc.save(dest)
print(str(dest))
print('Words:',len(' '.join(lines).split()),'Figures:',sum(x.startswith('![') for x in lines),'Tables:',len(doc.tables))
